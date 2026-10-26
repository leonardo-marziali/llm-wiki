"""Raw source files to cached markdown under .wiki/converted/."""

import csv
import importlib.util
import io
from dataclasses import dataclass, field
from pathlib import Path

from . import images, sheets
from .caps import Caps
from .errors import RefusedSource, UnsupportedFormat, WikiError

SUPPORTED = {
    ".md", ".txt", ".png", ".jpg", ".jpeg", ".gif", ".webp",
    ".pptx", ".xlsx", ".csv", ".docx", ".pdf",
}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".webp"}
NEVER_INGEST_SUFFIXES = (
    ".publishsettings", ".postman_environment.json", ".pem", ".key", ".pfx", ".p12", ".jks",
    ".keystore", ".env", ".kdbx", ".ppk",
)
MIN_PDF_CHARS_PER_PAGE = 200
CONVERTED_DIR = Path(".wiki") / "converted"


def never_ingest(filename: str) -> bool:
    return filename.lower().endswith(NEVER_INGEST_SUFFIXES)


@dataclass
class Conversion:
    converter: str
    outputs: list[str]
    chars: int
    parts: int
    flags: list[str] = field(default_factory=list)


def _header(source: str, sha: str, converter: str, part: str | None = None) -> str:
    part_line = f"part: {part}\n" if part else ""
    return f"---\nsource: {source}\nsha256: {sha}\nconverter: {converter}\n{part_line}---\n\n"


def _read_text(path: Path) -> str:
    data = path.read_bytes()
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("latin-1", errors="replace")


def _markitdown(path: Path) -> str:
    from markitdown import MarkItDown

    return MarkItDown().convert(str(path)).text_content


def _pptx_notes(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    blocks = []
    for i, slide in enumerate(prs.slides, 1):
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text.strip()
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        if notes:
            blocks.append(f"## Slide {i}: {title}\n\n{notes}\n")
    return "\n".join(blocks)


def _trim(rows: list[list[str]]) -> list[list[str]]:
    """Drop trailing all-empty rows and columns (formatted-but-empty cells)."""
    while rows and not any(c.strip() for c in rows[-1]):
        rows.pop()
    width = 0
    for row in rows:
        for i in range(len(row) - 1, -1, -1):
            if row[i].strip():
                width = max(width, i + 1)
                break
    return [r[:width] + [""] * (width - len(r)) for r in rows]


def _xlsx_rows(path: Path) -> dict[str, list[list[str]]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return {
            ws.title: _trim([["" if c is None else str(c) for c in row] for row in ws.iter_rows(values_only=True)])
            for ws in wb.worksheets
        }
    finally:
        wb.close()


def _csv_rows(path: Path) -> dict[str, list[list[str]]]:
    text = _read_text(path)
    return {path.stem: [list(r) for r in csv.reader(io.StringIO(text, newline=""))]}


def _pdf_pages(path: Path) -> int:
    from pypdf import PdfReader

    return max(len(PdfReader(str(path)).pages), 1)


def _pypdf_text(path: Path) -> str:
    from pypdf import PdfReader

    return "\n\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)


def _docling_text(path: Path) -> str:
    from docling.document_converter import DocumentConverter

    return DocumentConverter().convert(str(path)).document.export_to_markdown()


def _pdf(path: Path) -> tuple[str, str, list[str]]:
    pages = _pdf_pages(path)
    best_name, best_text = "markitdown", ""
    attempts = [("markitdown", _markitdown), ("pypdf", _pypdf_text)]
    if importlib.util.find_spec("docling") is not None:
        attempts.append(("docling", _docling_text))
    failed = 0
    for name, fn in attempts:
        try:
            text = fn(path)
        except Exception:
            failed += 1
            continue
        if len(text.strip()) > len(best_text.strip()):
            best_name, best_text = name, text
        if len(text.strip()) / pages >= MIN_PDF_CHARS_PER_PAGE:
            return name, text, []
    if failed == len(attempts):
        raise RefusedSource(f"{path.name}: could not be read (all PDF extractors failed)")
    flags = ["needs-vision"] if len(best_text.strip()) / pages < MIN_PDF_CHARS_PER_PAGE else []
    return best_name, best_text, flags


def _split_section(section: str, limit: int) -> list[str]:
    if len(section) <= limit:
        return [section]
    out: list[str] = []
    cur = ""
    for para in section.split("\n\n"):
        piece = para + "\n\n"
        while len(piece) > limit:  # a single paragraph over the limit: hard cut
            if cur:
                out.append(cur)
                cur = ""
            out.append(piece[:limit])
            piece = piece[limit:]
        if len(cur) + len(piece) > limit:
            out.append(cur)
            cur = ""
        cur += piece
    if cur:
        out.append(cur)
    return out


def _split(text: str, limit: int) -> list[str]:
    sections: list[str] = []
    cur: list[str] = []
    for line in text.splitlines(keepends=True):
        if (line.startswith("# ") or line.startswith("## ")) and cur:
            sections.append("".join(cur))
            cur = []
        cur.append(line)
    if cur:
        sections.append("".join(cur))
    parts: list[str] = []
    cur_part = ""
    for section in sections:
        for piece in _split_section(section, limit):
            if cur_part and len(cur_part) + len(piece) > limit:
                parts.append(cur_part)
                cur_part = ""
            cur_part += piece
    if cur_part:
        parts.append(cur_part)
    return parts


def convert(root: Path, raw_path: Path, sha: str, caps: Caps) -> Conversion:
    suffix = raw_path.suffix.lower()
    if never_ingest(raw_path.name):
        raise RefusedSource(f"{raw_path.name} looks like a secret/credential file and is never ingested")
    if suffix not in SUPPORTED:
        raise UnsupportedFormat(
            f"{suffix or raw_path.name} is not supported; supported: {', '.join(sorted(SUPPORTED))}"
        )

    out_dir = root / CONVERTED_DIR
    out_dir.mkdir(parents=True, exist_ok=True)
    try:
        source = "raw/" + raw_path.resolve().relative_to((root / "raw").resolve()).as_posix()
    except ValueError:
        source = "raw/" + raw_path.name

    flags: list[str] = []
    extra: list[Path] = []

    try:
        if suffix in (".md", ".txt"):
            converter, body = "text", _read_text(raw_path)
        elif suffix == ".docx":
            converter, body = "markitdown", _markitdown(raw_path)
        elif suffix == ".pptx":
            converter, body = "markitdown+pptx-notes", _markitdown(raw_path)
            notes = _pptx_notes(raw_path)
            if notes:
                notes_path = out_dir / f"{sha}.notes.md"
                notes_path.write_text(
                    _header(source, sha, converter) + notes, encoding="utf-8"
                )
                extra.append(notes_path)
        elif suffix in (".xlsx", ".csv"):
            converter = "sheets"
            rows = _xlsx_rows(raw_path) if suffix == ".xlsx" else _csv_rows(raw_path)
            body, csvs = sheets.convert_sheets(rows, out_dir, sha)
            if csvs:
                flags.append("large-sheet")
                extra.extend(csvs)
        elif suffix == ".pdf":
            converter, body, flags = _pdf(raw_path)
        else:
            info = images.triage(raw_path)
            if info.reason:
                raise RefusedSource(info.reason)
            converter = "image"
            body = f"Image: {info.width}x{info.height} px. Content not extracted; needs vision.\n"
            flags.append("needs-vision")
    except WikiError:
        raise
    except Exception as exc:
        raise RefusedSource(f"{raw_path.name}: could not be read ({type(exc).__name__})") from exc

    primary = out_dir / f"{sha}.md"
    text = _header(source, sha, converter) + body
    primary.write_text(text, encoding="utf-8")
    outputs = [primary]

    parts = 1
    if len(text) > caps.split_chars:
        head = _header(source, sha, converter)
        chunks = _split(body, max(caps.split_chars - len(head) - 32, 1))
        parts = len(chunks)
        for n, chunk in enumerate(chunks, 1):
            p = out_dir / f"{sha}.p{n}.md"
            p.write_text(_header(source, sha, converter, f"{n}/{parts}") + chunk, encoding="utf-8")
            outputs.append(p)
    outputs.extend(extra)

    return Conversion(
        converter=converter,
        outputs=[p.relative_to(root).as_posix() for p in outputs],
        chars=len(text),
        parts=parts,
        flags=flags,
    )
