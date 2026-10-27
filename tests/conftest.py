import json
from pathlib import Path

import pytest


@pytest.fixture(autouse=True)
def _isolated_env(tmp_path, monkeypatch):
    monkeypatch.setenv("GIT_AUTHOR_NAME", "test")
    monkeypatch.setenv("GIT_AUTHOR_EMAIL", "test@example.com")
    monkeypatch.setenv("GIT_COMMITTER_NAME", "test")
    monkeypatch.setenv("GIT_COMMITTER_EMAIL", "test@example.com")
    monkeypatch.setenv(
        "GIT_CEILING_DIRECTORIES", str(tmp_path.parent)
    )  # never see a repo above tmp_path
    monkeypatch.setenv("LLM_WIKI_HOME", str(tmp_path / "home"))
    monkeypatch.delenv("LLM_WIKI_PATH", raising=False)
    # A missing binary by default: no test can reach the real, paid claude. fake_claude
    # overrides it.
    monkeypatch.setenv("LLM_WIKI_CLAUDE", str(tmp_path / "no-claude"))


@pytest.fixture
def make_docx():
    def make(path, headings_and_paras):
        """Items are ("h1"|"h2"|"p", text)."""
        import docx

        doc = docx.Document()
        for kind, text in headings_and_paras:
            if kind == "p":
                doc.add_paragraph(text)
            else:
                doc.add_heading(text, level=int(kind[1:]))
        doc.save(str(path))
        return path

    return make


@pytest.fixture
def make_pptx():
    def make(path, slides):
        """slides = [(title, bullets, notes)]; notes may be empty."""
        from pptx import Presentation
        from pptx.shapes.autoshape import Shape

        prs = Presentation()
        for title, bullets, notes in slides:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            title_shape, body = slide.shapes.title, slide.placeholders[1]
            assert title_shape is not None and isinstance(body, Shape)
            title_shape.text = title
            body.text_frame.text = "\n".join(bullets)
            if notes:
                frame = slide.notes_slide.notes_text_frame
                assert frame is not None
                frame.text = notes
        prs.save(str(path))
        return path

    return make


@pytest.fixture
def make_xlsx():
    def make(path, sheets):
        import openpyxl
        from openpyxl.worksheet.worksheet import Worksheet

        wb = openpyxl.Workbook()
        default_sheet = wb.active
        assert isinstance(default_sheet, Worksheet)
        wb.remove(default_sheet)
        for name, rows in sheets.items():
            ws = wb.create_sheet(name)
            for row in rows:
                ws.append(row)
        wb.save(str(path))
        return path

    return make


@pytest.fixture
def make_csv():
    def make(path, rows, encoding="utf-8"):
        import csv

        with open(path, "w", newline="", encoding=encoding) as fh:
            csv.writer(fh).writerows(rows)
        return path

    return make


@pytest.fixture
def make_text_pdf():
    def make(path, pages):
        """Minimal hand-written PDF: one Helvetica Tj text object per page."""
        n = len(pages)
        objs = {1: b"<< /Type /Catalog /Pages 2 0 R >>"}
        kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(n))
        objs[2] = f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode()
        font = 3 + 2 * n
        for i, text in enumerate(pages):
            esc = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            stream = f"BT /F1 12 Tf 72 720 Td ({esc}) Tj ET".encode()
            objs[3 + 2 * i] = (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                f"/Resources << /Font << /F1 {font} 0 R >> >> "
                f"/Contents {4 + 2 * i} 0 R >>"
            ).encode()
            objs[4 + 2 * i] = (
                f"<< /Length {len(stream)} >>\nstream\n".encode()
                + stream
                + b"\nendstream"
            )
        objs[font] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
        out = bytearray(b"%PDF-1.4\n")
        offsets = {}
        for num in sorted(objs):
            offsets[num] = len(out)
            out += f"{num} 0 obj\n".encode() + objs[num] + b"\nendobj\n"
        xref = len(out)
        size = max(objs) + 1
        out += f"xref\n0 {size}\n0000000000 65535 f \n".encode()
        for num in range(1, size):
            out += f"{offsets[num]:010d} 00000 n \n".encode()
        out += (
            f"trailer\n<< /Size {size} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
        ).encode()
        path.write_bytes(bytes(out))
        return path

    return make


@pytest.fixture
def make_png():
    def make(path, size=(64, 64), rects=((8, 8, 40, 40),)):
        from PIL import Image, ImageDraw

        img = Image.new("RGB", size, "white")
        draw = ImageDraw.Draw(img)
        for rect in rects:
            draw.rectangle(rect, fill="black")
        img.save(str(path))
        return path

    return make


@pytest.fixture
def make_scanned_pdf():
    def make(path):
        from PIL import Image, ImageDraw

        img = Image.new("RGB", (400, 400), "white")
        ImageDraw.Draw(img).rectangle((50, 50, 300, 300), fill="gray")
        img.save(str(path), "PDF")
        return path

    return make


RAW_REL = "raw/docs/guide.md"
PAGE_REL = "conventions/git-flow.md"


def page_text(
    id_: str,
    title: str,
    body: str,
    *,
    type_: str = "convention",
    status: str = "as-is",
    sources: list[tuple[str, str]] | None = None,
    extra: str = "",
    last_verified: str = "2026-01-01",
    review_after: str = "2099-01-01",
) -> str:
    """Frontmatter + body for a page; ``sources`` is a list of (path, sha256)."""
    if sources:
        src = "sources:\n" + "".join(
            f"- path: {p}\n  sha256: {s}\n  locator: section 1\n" for p, s in sources
        )
    else:
        src = "sources: []\n"
    return (
        "---\n"
        f"id: {id_}\ntitle: {title}\ntype: {type_}\nstatus: {status}\n{extra}"
        f"summary: {title} summary.\n{src}"
        "confidence: high\nreview_status: draft\n"
        f"last_verified: {last_verified}\nreview_after: {review_after}\n---\n{body}"
    )


@pytest.fixture
def make_wiki():
    from llm_wiki.config import Profile

    def make(tmp_path, profile: Profile = "software"):
        from datetime import date

        from llm_wiki.config import WikiConfig, dump_config, load_profile
        from llm_wiki.index import build_indexes
        from llm_wiki.manifest import Manifest, SourceEntry, sha256_file

        root = tmp_path / "w"
        cfg = WikiConfig(
            name="Acme",
            description="Acme wiki.",
            profile=profile,
            sections=load_profile(profile),
        )
        root.mkdir(parents=True)
        (root / "wiki.toml").write_text(dump_config(cfg), encoding="utf-8")
        for s in cfg.sections:
            (root / "wiki" / s.name).mkdir(parents=True)
        raw = root / RAW_REL
        raw.parent.mkdir(parents=True)
        raw.write_text("# Guide\n\nAlways branch from main.\n", encoding="utf-8")
        sha = sha256_file(raw)
        m = Manifest()
        m.record(
            SourceEntry(
                path=RAW_REL,
                sha256=sha,
                size=raw.stat().st_size,
                mime="text/markdown",
                category="docs",
                first_seen=date(2026, 1, 1),
                status="ingested",
            )
        )
        m.save(root)
        wiki = root / "wiki"
        if profile == "software":
            (wiki / PAGE_REL).write_text(
                page_text(
                    "conventions.git-flow",
                    "Git flow",
                    "# Git flow\n\n**TL;DR** Branch from main.\n\n"
                    "We branch from main for every change.[^1]\n\n"
                    f"[^1]: {RAW_REL}\n",
                    sources=[(RAW_REL, sha)],
                ),
                encoding="utf-8",
            )
            overview_body = (
                "# Overview\n\n**TL;DR** Start here.\n\n"
                "- [Git flow](conventions/git-flow.md)\n"
            )
        else:
            overview_body = "# Overview\n\n**TL;DR** Start here.\n"
        (wiki / "overview.md").write_text(
            page_text("overview", "Overview", overview_body, type_="overview"),
            encoding="utf-8",
        )
        (wiki / "glossary.md").write_text(
            page_text(
                "glossary",
                "Glossary",
                "# Glossary\n\n**TL;DR** Terms.\n",
                type_="glossary",
            ),
            encoding="utf-8",
        )
        (wiki / "open-questions.md").write_text("# Open questions\n", encoding="utf-8")
        build_indexes(root, cfg)
        (wiki / "log.md").write_text(
            "# Log\n\n## [2026-10-26] init | created wiki\n", encoding="utf-8"
        )
        return root

    return make


def ok(result="done", structured=None, cost=0.01):
    """A successful `claude -p --output-format json` payload."""
    return {
        "type": "result",
        "subtype": "success",
        "is_error": False,
        "result": result,
        "structured_output": structured,
        "total_cost_usd": cost,
        "duration_ms": 5,
        "session_id": "s1",
    }


class FakeClaude:
    def __init__(self, root):
        self.queue_file = root / "queue.json"
        self.log_file = root / "log.jsonl"
        self.queue()

    def queue(self, *responses):
        self.queue_file.write_text(json.dumps(list(responses)), encoding="utf-8")

    def calls(self) -> list[dict]:
        if not self.log_file.exists():
            return []
        return [
            json.loads(line)
            for line in self.log_file.read_text(encoding="utf-8").splitlines()
        ]


@pytest.fixture
def fake_claude(tmp_path, monkeypatch):
    from llm_wiki.claude import REQUIRED_FLAGS

    fake = FakeClaude(tmp_path)
    monkeypatch.setenv(
        "LLM_WIKI_CLAUDE", str(Path(__file__).parent / "fake_claude" / "claude")
    )
    monkeypatch.setenv("FAKE_CLAUDE_QUEUE", str(fake.queue_file))
    monkeypatch.setenv("FAKE_CLAUDE_LOG", str(fake.log_file))
    monkeypatch.setenv(
        "FAKE_CLAUDE_HELP", "\n".join(REQUIRED_FLAGS)
    )  # one source of truth
    return fake
